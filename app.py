from fileinput import filename
import os
from annotated_types import doc
from flask import Flask, render_template, request, send_file
from groq import Groq
from docx import Document
from dotenv import load_dotenv

# ================================
# LOAD ENV VARIABLES
# ================================
load_dotenv()

app = Flask(__name__)

# ================================
# CONFIGURATION
# ================================
MODEL_NAME = "llama-3.3-70b-versatile"

api_key = os.getenv("GROQ_API_KEY")
if not api_key:
    raise ValueError("❌ GROQ_API_KEY not found in .env file")

client = Groq(api_key=api_key)

# ================================
# MASTER PROMPT
# ================================
TEMPLATE_PROMPT = """
You are a professional legal drafting assistant.

The user will ONLY specify the type of legal document required.
Your task is to generate a COMPLETE, professionally structured legal document TEMPLATE.

Rules:
- Do NOT ask for user details
- Do NOT include real names, dates, or locations
- Use ONLY editable placeholders inside square brackets
- Use formal legal language
- Ensure the document is reusable and industry-standard
- Include all standard clauses relevant to the document type
- Output ONLY the legal document template

User Instruction:
{instruction}
"""

# ================================
# ROUTES
# ================================
@app.route("/", methods=["GET", "POST"])
def index():
    if request.method == "POST":
        instruction = request.form.get("instruction")

        if not instruction:
            return "Instruction cannot be empty", 400

        prompt = TEMPLATE_PROMPT.format(instruction=instruction)

        response = client.chat.completions.create(
            model=MODEL_NAME,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.2
        )

        output_text = response.choices[0].message.content.strip()

        # Save DOCX
        doc = Document()
        for line in output_text.split("\n"):
            doc.add_paragraph(line)

        #filename = instruction.lower().replace(" ", "_") + ".docx"
        #doc.save(filename)

        filename = "generated_docs/" + instruction.lower().replace(" ", "_") + ".docx"
        doc.save(filename)



        return send_file(filename, as_attachment=True)

    return render_template("index.html")


# ================================
# RUN
# ================================
if __name__ == "__main__":
    app.run(debug=True)
